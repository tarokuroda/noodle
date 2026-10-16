"""Backend smoke tests. Run from the repo root: python -m pytest backend/tests -q"""

import base64
import hashlib
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from noodle.db import connect  # noqa: E402
from noodle.moodle import clean, mask, normalize_site, parse_app_link  # noqa: E402


def test_clean_makes_safe_windows_names():
    assert clean('a:b/c?"d') == "a_b_c__d"
    assert clean("CON.txt") == "_CON.txt"
    assert clean("trailing dot. ") == "trailing dot"


def test_normalize_site():
    assert normalize_site("moodle.ucl.ac.uk/my/") == "https://moodle.ucl.ac.uk"
    assert normalize_site("https://moodle.ucl.ac.uk/login/index.php?x=1") == "https://moodle.ucl.ac.uk"


def test_parse_app_link_and_mask():
    site, passport, token = "https://moodle.example", "12345", "a" * 32
    sig = hashlib.md5((site + passport).encode()).hexdigest()
    link = "moodlemobile://token=" + base64.b64encode(f"{sig}:::{token}".encode()).decode()
    assert parse_app_link(link, site, passport) == (token, True)
    assert parse_app_link(token, site, passport) == (token, True)
    assert token not in mask(token)


def test_search_finds_course_material():
    from noodle.search import search
    course = connect().execute("SELECT id FROM courses ORDER BY synced_at DESC").fetchone()
    if course is None:
        return  # nothing synced on this machine
    hits = search("nozzle flow", course["id"])
    assert hits and all({"file_id", "page", "name", "snippet"} <= set(h) for h in hits)


def test_wikilinks_parse_and_resolve():
    from noodle import notes
    found = notes.links("See [[Lecture 3.pdf#page=12|p.12]] and [[My note]].")
    assert found == [("Lecture 3.pdf", "page=12", "p.12"), ("My note", None, None)]
    course = connect().execute("SELECT id FROM courses ORDER BY synced_at DESC").fetchone()
    if course is None:
        return
    targets = notes.link_targets(course["id"])
    pdf = next(t for t in targets if t["kind"] == "pdf")
    assert notes.resolve(pdf["link"], course["id"], targets) == ("file", pdf["id"])


def test_slide_text_fallback(tmp_path):
    from pptx import Presentation
    from noodle.convert import slide_texts
    deck = Presentation()
    for title in ("Isentropic flow", "Normal shocks"):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"About {title.lower()}"
    path = tmp_path / "deck.pptx"
    deck.save(path)
    pages = slide_texts(path)
    assert [n for n, _ in pages] == [1, 2]
    assert "Isentropic flow" in pages[0][1] and "normal shocks" in pages[1][1]


def test_viewer_and_twins():
    from noodle.index import twin_key, viewer_for
    row = lambda kind, pdf=None, pages=None: {"kind": kind, "pdf": pdf, "pages": pages}  # noqa: E731
    assert viewer_for(row("pdf")) == "pdf" and viewer_for(row("page")) == "page"
    assert viewer_for(row("office", pdf=".noodle/converted/x.pdf")) == "pdf"
    assert viewer_for(row("office", pages=12)) == "page"
    assert viewer_for(row("office")) == "app" and viewer_for(row("media")) == "app"
    assert twin_key("c/Week 4/Lecture4.pptx") == twin_key("c/Week 4/LECTURE4.pdf")
    assert twin_key("c/Week 4/Lecture4.pptx") != twin_key("c/Week 5/Lecture4.pdf")


def test_timeline_merges_calendar_events():
    from noodle.timeline import merge_events
    now = 1_800_000_000
    action = [{"id": 1, "name": "Coursework 1 is due", "eventtype": "due", "modulename": "assign", "instance": 77,
               "timesort": now + 86400, "url": "https://m.example/mod/assign/view.php?id=77", "overdue": False,
               "course": {"id": 5, "shortname": "ELEC0134_26-27"}},
              {"id": 3, "name": "Quiz closes", "eventtype": "close", "modulename": "quiz", "instance": 80,
               "timesort": now - 86400, "url": "https://m.example/mod/quiz/view.php?id=80", "overdue": True,
               "course": {"id": 5, "shortname": "ELEC0134_26-27"}}]
    raw = [{"id": 1, "name": "Coursework 1 is due", "eventtype": "due", "modulename": "assign", "instance": 77,
            "timestart": now + 86400, "courseid": 5},
           {"id": 2, "name": "Lab report is due", "eventtype": "due", "modulename": "assign", "instance": 78,
            "timestart": now + 2 * 86400, "courseid": 5},
           {"id": 4, "name": "Marking", "eventtype": "gradingdue", "modulename": "assign", "instance": 79,
            "timestart": now, "courseid": 5}]
    statuses = {"77": {"state": "done", "text": "Submitted"}}
    items = {i["id"]: i for i in merge_events(action, raw, statuses, {5: "ELEC0134"}, "https://m.example", now)}
    assert set(items) == {"ev1", "ev2", "ev3"}  # merged by id, teacher-only marking deadline dropped
    assert items["ev1"]["status"]["text"] == "Submitted" and items["ev1"]["course"] == "ELEC0134"
    assert items["ev2"]["url"].startswith("https://m.example/calendar/") and items["ev2"]["kind"] == "deadline"
    assert items["ev3"]["overdue"] and not items["ev1"]["overdue"]


def test_timeline_offline_keeps_cache(tmp_path, monkeypatch):
    import requests
    from noodle import account, timeline
    monkeypatch.setattr(timeline, "CACHE_PATH", tmp_path / "timeline.json")
    timeline.save_json(timeline.CACHE_PATH, {"fetched_at": 1.0, "site": "https://m.example", "courses": {},
                                             "action": [], "raw": [], "statuses": {}})

    class Down:
        site, token = "https://m.example", "x" * 32

        def call(self, *a, **k):
            raise requests.ConnectionError("offline")

    monkeypatch.setattr(account, "status", lambda: {"logged_in": True, "userid": 1})
    out = timeline.refresh(Down(), force=True)
    assert out["offline"] and out["fetched_at"] == 1.0
