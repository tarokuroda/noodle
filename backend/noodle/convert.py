"""Slides and Word docs -> PDF, so they open in the viewer with highlights, search, Ask and the graph.

PowerPoint and Word do the conversion through COM, one file at a time, each in a child process with a timeout
(an Office dialog that never closes can't hang the sync). The PDFs are cached in .noodle/converted/ by the
source file's sha1, so a deck is only converted again when it changes. If Office is missing or a file fails,
slides fall back to their python-pptx text, one page per slide, so search and Ask still find them.

    python -m noodle.convert SRC OUT      (from backend/) the child process: convert one file
"""

import csv
import subprocess
import sys
import time
from collections import Counter
from pathlib import Path

from .paths import CONVERTED_DIR, ROOT

CONVERTIBLE = ("pptx", "ppt", "docx", "doc")
WORD = (".doc", ".docx")
TIMEOUT = 300  # seconds; the 50 MB Lecture 1 deck takes about a minute
OFFICE_EXES = ("POWERPNT.EXE", "WINWORD.EXE")


def app_name(ext):
    ext = (ext or "").lower()
    return "Word" if ext in ("doc", "docx") else "PowerPoint" if ext in ("ppt", "pptx") else None


# ---------- the parent side: run one conversion, never more than TIMEOUT ----------

def _office_pids():
    pids = set()
    for exe in OFFICE_EXES:
        out = subprocess.run(["tasklist", "/FI", f"IMAGENAME eq {exe}", "/FO", "CSV", "/NH"],
                             capture_output=True, text=True).stdout
        pids |= {int(row[1]) for row in csv.reader(out.splitlines()) if len(row) > 1 and row[1].isdigit()}
    return pids


def to_pdf(src, out, timeout=TIMEOUT):
    """Convert src to out (a .pdf). Returns None, or a short error message."""
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(out.stem + ".tmp.pdf")
    before = _office_pids()
    try:
        r = subprocess.run([sys.executable, "-m", "noodle.convert", str(src), str(tmp)], cwd=Path(__file__).parents[1],
                           capture_output=True, text=True, timeout=timeout)
    except subprocess.TimeoutExpired:
        for pid in _office_pids() - before:  # only the Office we started; your own PowerPoint stays open
            subprocess.run(["taskkill", "/PID", str(pid), "/F"], capture_output=True)
        tmp.unlink(missing_ok=True)
        return f"timed out after {timeout} s"
    if r.returncode != 0 or not tmp.exists() or tmp.stat().st_size == 0:
        tmp.unlink(missing_ok=True)
        lines = (r.stderr or r.stdout or "no output").strip().splitlines()
        return lines[-1][:300] if lines else "no output"
    tmp.replace(out)
    return None


def convert_course(course_id=None, log=print, progress=None):
    """Convert the course's slides and Word docs that have no up-to-date PDF yet (twins of a PDF are skipped)."""
    from .db import connect
    from .index import file_hash
    conn = connect()
    marks = ",".join("?" * len(CONVERTIBLE))
    rows = conn.execute(
        f"SELECT id, path, name, pdf, indexed_mtime FROM files WHERE status='ok' AND kind='office' "
        f"AND twin_id IS NULL AND ext IN ({marks})" + (" AND course_id=?" if course_id else ""),
        (*CONVERTIBLE, *((course_id,) if course_id else ()))).fetchall()
    stats = Counter()
    for n, row in enumerate(rows):
        src = ROOT / row["path"]
        if not src.exists():
            continue
        if row["indexed_mtime"] == src.stat().st_mtime and (not row["pdf"] or (ROOT / row["pdf"]).exists()):
            continue  # this version is already converted (or already fell back to its text)
        h = file_hash(src)
        out = CONVERTED_DIR / f"{h}.pdf"
        failed = CONVERTED_DIR / f"{h}.failed"  # delete it to try that file again
        if not out.exists() and not failed.exists():
            if progress:
                progress("Converting slides", n + 1, len(rows))
            t = time.time()
            err = to_pdf(src, out)
            if err:
                failed.write_text(err, encoding="utf-8")
                stats["failed"] += 1
                log(f"  ! Couldn't convert {row['path']}: {err} (using its text instead)")
            else:
                stats["converted"] += 1
                log(f"  converted {row['name']} ({time.time() - t:.0f} s)")
        pdf = out.relative_to(ROOT).as_posix() if out.exists() else None
        conn.execute("UPDATE files SET hash=?, pdf=?, indexed_mtime=NULL WHERE id=?", (h, pdf, row["id"]))
        conn.commit()
    return stats


def slide_texts(path):
    """[(slide_number, text)] from a .pptx, for when PowerPoint can't convert it."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def texts(shape):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in shape.shapes:
                yield from texts(s)
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield shape.text_frame.text
        if getattr(shape, "has_table", False) and shape.has_table:
            for r in shape.table.rows:
                yield " | ".join(c.text for c in r.cells)

    return [(i + 1, "\n".join(t for s in slide.shapes for t in texts(s) if t.strip()))
            for i, slide in enumerate(Presentation(str(path)).slides)]


# ---------- the child side: one file through COM ----------

def _powerpoint(src, out):
    import win32com.client
    try:
        win32com.client.GetActiveObject("PowerPoint.Application")
        was_running = True
    except Exception:  # noqa: BLE001 - "not running" comes back as a COM error
        was_running = False
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(str(src), -1, 0, 0)  # ReadOnly, not Untitled, no window
    try:
        pres.SaveAs(str(out), 32)  # ppSaveAsPDF
    finally:
        pres.Close()
        if not was_running and app.Presentations.Count == 0:
            app.Quit()  # never close a PowerPoint you had open


def _word(src, out):
    import win32com.client
    app = win32com.client.DispatchEx("Word.Application")  # a private Word, separate from yours
    app.Visible = False
    app.DisplayAlerts = 0
    try:
        doc = app.Documents.Open(str(src), False, True, False)  # no conversion prompt, read-only, not in recents
        try:
            doc.ExportAsFixedFormat(str(out), 17)  # wdExportFormatPDF
        finally:
            doc.Close(0)
    finally:
        app.Quit()


def main(argv):
    import pythoncom
    src, out = Path(argv[1]).resolve(), Path(argv[2]).resolve()
    pythoncom.CoInitialize()
    try:
        (_word if src.suffix.lower() in WORD else _powerpoint)(src, out)
    finally:
        pythoncom.CoUninitialize()


if __name__ == "__main__":
    main(sys.argv)
